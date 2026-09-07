"""
Module: extractor.py
Description: Trích xuất nội dung văn bản, speaker notes và hình ảnh độ phân giải cao từ file PowerPoint (.pptx).
"""

import os
import sys
import json
from pathlib import Path
from typing import List, Dict, Any
from pptx import Presentation
from PIL import Image, ImageDraw

try:
    from . import config
except (ImportError, ValueError):
    import config

def extract_text_and_notes(pptx_path: Path) -> List[Dict[str, Any]]:
    """Trích xuất tiêu đề slide, các gạch đầu dòng, bảng biểu và Speaker Notes."""
    prs = Presentation(str(pptx_path))
    slides_data = []

    for idx, slide in enumerate(prs.slides, start=1):
        title = ""
        bullet_points = []
        notes = ""

        # 1. Trích xuất Text & Tables
        for shape in slide.shapes:
            if shape.has_text_frame:
                text_frame = shape.text_frame
                text = text_frame.text.strip()
                if not text:
                    continue

                if shape == slide.shapes.title or (not title and len(text) < 120 and "\n" not in text):
                    title = text
                else:
                    for paragraph in text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text and p_text not in bullet_points:
                            bullet_points.append(p_text)

            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        bullet_points.append(f"Bảng dữ liệu: {row_text}")

        # 2. Trích xuất Speaker Notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if "Click to add notes" in notes:
                notes = ""

        if not title:
            title = f"Slide {idx}"

        slides_data.append({
            "slide_index": idx,
            "title": title,
            "bullet_points": bullet_points,
            "speaker_notes": notes,
            "image_path": str(config.IMAGES_DIR / f"slide_{idx:03d}.png")
        })

    return slides_data


def export_slides_as_images(pptx_path: Path, output_dir: Path) -> bool:
    """Xuất tất cả slide thành ảnh PNG 1080p bằng PowerPoint COM (hoặc Fallback)."""
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Xóa sạch các ảnh slide cũ để không bị sót slide từ file bài giảng trước
    for old_img in output_dir.glob("slide_*.png"):
        try: old_img.unlink()
        except Exception: pass

    abs_pptx = str(pptx_path.resolve())

    # 1. Thử Office COM
    try:
        import win32com.client
        import pythoncom

        pythoncom.CoInitialize()
        powerpoint = win32com.client.Dispatch("PowerPoint.Application")
        try:
            powerpoint.Visible = 1
        except Exception:
            pass

        deck = powerpoint.Presentations.Open(abs_pptx, ReadOnly=True, Untitled=False, WithWindow=False)
        for i, slide in enumerate(deck.Slides, start=1):
            image_dest = output_dir / f"slide_{i:03d}.png"
            slide.Export(str(image_dest.resolve()), "PNG", config.VIDEO_WIDTH, config.VIDEO_HEIGHT)

        deck.Close()
        powerpoint.Quit()
        pythoncom.CoUninitialize()
        return True
    except Exception as e:
        print(f"[!] PowerPoint COM Warning: {e}. Sử dụng trình tạo slide card dự phòng...")

    # 2. Fallback vẽ slide card nếu không có PowerPoint COM
    prs = Presentation(str(pptx_path))
    for i, slide_data in enumerate(prs.slides, start=1):
        image_dest = output_dir / f"slide_{i:03d}.png"
        img = Image.new("RGB", (config.VIDEO_WIDTH, config.VIDEO_HEIGHT), color=(15, 23, 42))
        draw = ImageDraw.Draw(img)

        # Banner
        draw.rectangle([(0, 0), (config.VIDEO_WIDTH, 120)], fill=(30, 41, 59))
        draw.rectangle([(0, 116), (config.VIDEO_WIDTH, 120)], fill=(14, 165, 233))

        # Card
        draw.rounded_rectangle(
            [(80, 180), (config.VIDEO_WIDTH - 80, config.VIDEO_HEIGHT - 80)],
            radius=20, fill=(24, 32, 47), outline=(51, 65, 85), width=2
        )

        title_text = f"Slide {i}"
        for shape in slide_data.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                title_text = shape.text_frame.text.strip().split("\n")[0][:80]
                break

        draw.text((100, 45), f"AI LECTURE STUDIO  |  SLIDE {i}", fill=(148, 163, 184))
        draw.text((120, 220), title_text, fill=(255, 255, 255))
        img.save(image_dest)

    return True


def run_extractor(pptx_file: Path) -> List[Dict[str, Any]]:
    """Hàm thực thi chính cho Bước 1: Trích xuất Slide"""
    if not pptx_file.exists():
        raise FileNotFoundError(f"Không tìm thấy file: {pptx_file}")

    slides_data = extract_text_and_notes(pptx_file)
    export_slides_as_images(pptx_file, config.IMAGES_DIR)

    output_json = config.WORKSPACE_DIR / "extracted_data.json"
    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(slides_data, f, ensure_ascii=False, indent=2)

    return slides_data

run_step1 = run_extractor
