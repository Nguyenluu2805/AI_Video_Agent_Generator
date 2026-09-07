import sys
import io
# Ensure UTF-8 output on Windows console
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_demo_pptx(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # ----------------------------------------------------
    # Slide 1: Title Slide
    # ----------------------------------------------------
    slide_layout = prs.slide_layouts[6] # Blank
    slide1 = prs.slides.add_slide(slide_layout)
    
    bg = slide1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg.line.color.rgb = RGBColor(15, 23, 42)
    
    txBox = slide1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.0))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "XÂY DỰNG AI AGENT TỰ ĐỘNG HÓA"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Quy trình chuyển đổi bài giảng PowerPoint thành Video hoàn chỉnh"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(226, 232, 240)
    p2.alignment = PP_ALIGN.CENTER
    
    notes_slide = slide1.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Chào mừng tất cả các bạn đến với khóa học chuyên sâu về AI Agent. Hôm nay tôi sẽ hướng dẫn các bạn cách thức xây dựng một hệ thống tạo video bài giảng tự động từ slide PowerPoint."
    
    # ----------------------------------------------------
    # Slide 2: Kiến Trúc 5 Bước
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(slide_layout)
    bg2 = slide2.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg2.line.color.rgb = RGBColor(15, 23, 42)
    
    title_box = slide2.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.0))
    tf_title = title_box.text_frame
    p = tf_title.paragraphs[0]
    p.text = "Kiến Trúc Pipeline Xử Lý 5 Bước"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    
    content_box = slide2.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.8), Inches(4.5))
    tf_content = content_box.text_frame
    tf_content.word_wrap = True
    
    steps = [
        "1. Trích xuất dữ liệu: Sử dụng win32com và python-pptx để lấy ảnh slide 1080p và ghi chú.",
        "2. AI Viết kịch bản: LLM biến gạch đầu dòng thành lời giảng sư phạm truyền cảm.",
        "3. Sinh giọng nền: Edge-TTS tạo nhịp điệu và ngữ điệu tiếng Việt tự nhiên.",
        "4. Voice Cloning RVC: Sử dụng GPU RTX 3060 Ti để nhái lại chính xác âm sắc của giảng viên.",
        "5. Lắp ráp Video: FFmpeg tự động căn khớp thời gian ảnh và tiếng để xuất file MP4."
    ]
    
    for i, step_text in enumerate(steps):
        if i == 0:
            p = tf_content.paragraphs[0]
        else:
            p = tf_content.add_paragraph()
        p.text = step_text
        p.font.size = Pt(20)
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.space_after = Pt(14)
        
    notes2 = slide2.notes_slide.notes_text_frame
    notes2.text = "Toàn bộ hệ thống được chia thành 5 module độc lập. Điểm đặc biệt là ở bước 4, chúng ta tận dụng sức mạnh tính toán song song của card đồ họa RTX 3060 Ti để xử lý chuyển đổi giọng nói trong tích tắc."
    
    # ----------------------------------------------------
    # Slide 3: Lợi Ích Thực Tiễn
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(slide_layout)
    bg3 = slide3.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = RGBColor(15, 23, 42)
    bg3.line.color.rgb = RGBColor(15, 23, 42)
    
    title_box3 = slide3.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(1.0))
    p = title_box3.text_frame.paragraphs[0]
    p.text = "Lợi Ích Và Hiệu Quả Thực Tế"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(56, 189, 248)
    
    content_box3 = slide3.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.8), Inches(4.5))
    tf_content3 = content_box3.text_frame
    tf_content3.word_wrap = True
    
    benefits = [
        "Tiết kiệm 90% thời gian so với việc thu âm và chỉnh sửa video thủ công.",
        "Đồng bộ hóa hoàn hảo giữa lời giảng và nội dung hiển thị trên từng trang slide.",
        "Dễ dàng cập nhật nội dung bài giảng mà không cần phải setup lại phòng thu âm.",
        "Chi phí vận hành gần như bằng 0 nhờ tận dụng GPU cục bộ và mô hình mã nguồn mở."
    ]
    
    for i, ben in enumerate(benefits):
        if i == 0:
            p = tf_content3.paragraphs[0]
        else:
            p = tf_content3.add_paragraph()
        p.text = f"• {ben}"
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(241, 245, 249)
        p.space_after = Pt(18)
        
    notes3 = slide3.notes_slide.notes_text_frame
    notes3.text = "Như các bạn đã thấy, ứng dụng AI Agent vào giảng dạy sẽ giúp các thầy cô và các nhà sáng tạo nội dung nâng cao hiệu suất làm việc lên gấp nhiều lần. Cảm ơn các bạn đã theo dõi bài giảng!"
    
    prs.save(str(output_path))
    print(f"[OK] Đã tạo thành công file PowerPoint mẫu tại: {output_path}")

if __name__ == "__main__":
    out = Path("input/demo_lecture.pptx")
    create_demo_pptx(out)
